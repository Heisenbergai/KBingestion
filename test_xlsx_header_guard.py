"""
Regression guard for the xlsx extractors.

REPLACES the scratchpad guard from the structured-spreadsheet work (whose
SHA256 9631df5c… is recorded in 06 but whose script was lost with the
session). Same purpose, stronger property, and now committed so it cannot
go missing again.

THE PROPERTY THIS PROTECTS. extract_xlsx() is on the live path for every
upload AND every Google Drive sync, and its prose output becomes chunk text
-> embeddings. Changing it for ordinary spreadsheets would silently alter
retrieval for every future ingest. So banner-row detection must be
surgical:

  * a NORMAL sheet (header on row 1) must come out BYTE-IDENTICAL, and
  * a BANNER sheet (single merged title above the real header) must switch
    from ["TITLE","col_1",...] to the real column names.

Run:  python3 test_xlsx_header_guard.py
"""
import hashlib
import io

import openpyxl

import ingest


def _wb(rows_by_sheet: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in rows_by_sheet.items():
        ws = wb.create_sheet(title=name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Fixture A: ORDINARY sheets. Must never change. -----------------------
# Deliberately mirrors the original guard's shape: mixed types, a blank row,
# a second sheet.
NORMAL = _wb({
    "Revenue": [
        ("Quarter", "Revenue", "Closed", "Booked On"),
        ("Q1", 4200000, True, "2024-04-01"),
        ("Q2", 5100000.5, False, "2024-07-01"),
        (None, None, None, None),          # blank row
        ("Q3", 6000000, True, "2024-10-01"),
    ],
    "Headcount": [
        ("Team", "People"),
        ("Engineering", 42),
        ("Sales", 17),
    ],
})

# --- Fixture B: BANNER sheet. This is what P1-14 gets wrong. --------------
BANNER = _wb({
    "Budget Summary": [
        ("MASTER BUDGET SUMMARY", None, None, None),   # merged title banner
        ("Department", "Q1 Actual", "Q2 Actual", "Variance"),
        ("Engineering", 1850000, 1920000, 70000),
        ("Sales", 900000, 875000, -25000),
    ],
})

# --- Fixture C: banner-like but AMBIGUOUS. Must be left alone. ------------
# A genuine one-column sheet: first row has one cell, second row also has
# one cell. Treating row 1 as a banner here would eat a real header.
ONE_COL = _wb({
    "Notes": [
        ("Note",),
        ("Renewal due in March",),
        ("Escalate to legal",),
    ],
})

# --- Fixture D: TWO-ROW banner (title, then a blank/subtitle spacer) before
# the real header — added 2026-08-13 while de-risking Phase 0 of the
# dashboard-upgrade thread ahead of Tanmay reprocessing the real corpus.
# Several live "Executive Dashboard" sheets look like this shape (a merged
# title row, then a narrower subtitle/date row, then the real header).
TWO_ROW_BANNER = _wb({
    "Exec Dashboard": [
        ("FY2024-25 ENTERPRISE ANNUAL BUDGET DASHBOARD", None, None, None, None),
        ("As of Q2", None, None, None, None),
        ("Metric", "Target", "Actual", "Variance", "Status"),
        ("Revenue", 5000000, 5200000, 200000, "On track"),
        ("Headcount", 120, 118, -2, "Under"),
    ],
})

# --- Fixture E: banner deeper than the scan bound — must be left alone. ---
# Four single-cell rows before the real header, one more than
# _BANNER_SCAN_LIMIT (3) can reach. The fix must not partially resolve this;
# it should fall back to treating row 0 as the header, same as an unbounded
# single-column sheet, rather than guessing past its proven-safe limit.
DEEP_BANNER = _wb({
    "Too Deep": [
        ("TITLE", None, None),
        ("Subtitle", None, None),
        ("Prepared by Finance", None, None),
        ("Confidential — internal use only", None, None),
        ("Metric", "Value", "Notes"),
        ("Revenue", 1000000, "Q2"),
    ],
})

# --- Fixture F: a numeric-looking rank/ID column right under the real
# header — confirms header detection looks only at row WIDTH, never at
# whether the row's values happen to look numeric (a live "Sales
# Leaderboard" sheet has exactly this shape: Rank 1, 2, 3... under a real
# header, which the pre-fix code had folded into a fake "numeric" banner
# column).
RANKED = _wb({
    "Leaderboard": [
        ("Q2 INDIVIDUAL PERFORMANCE LEADERBOARD", None, None, None),
        ("Rank", "Rep", "Deals Closed", "Revenue"),
        (1, "A. Rao", 14, 980000),
        (2, "S. Iyer", 11, 845000),
    ],
})


def _digest(obj) -> str:
    return hashlib.sha256(repr(obj).encode()).hexdigest()


def report():
    out = {}

    prose_normal = ingest.extract_xlsx(NORMAL)
    out["normal_prose_sha256"] = hashlib.sha256(prose_normal.encode()).hexdigest()
    out["normal_prose_len"] = len(prose_normal)

    tables_normal = ingest.extract_xlsx_tables(NORMAL)
    out["normal_tables_sha256"] = _digest(tables_normal)
    out["normal_headers"] = [t["headers"] for t in tables_normal]

    prose_banner = ingest.extract_xlsx(BANNER)
    out["banner_prose_len"] = len(prose_banner)
    tables_banner = ingest.extract_xlsx_tables(BANNER)
    out["banner_headers"] = [t["headers"] for t in tables_banner]
    out["banner_numeric_cols"] = [t["numeric_columns"] for t in tables_banner]
    out["banner_first_row"] = [t["rows"][0] for t in tables_banner]

    tables_onecol = ingest.extract_xlsx_tables(ONE_COL)
    out["onecol_headers"] = [t["headers"] for t in tables_onecol]
    out["onecol_row_count"] = [t["row_count"] for t in tables_onecol]

    tables_tworow = ingest.extract_xlsx_tables(TWO_ROW_BANNER)
    out["tworow_headers"] = [t["headers"] for t in tables_tworow]
    out["tworow_numeric_cols"] = [t["numeric_columns"] for t in tables_tworow]

    tables_deep = ingest.extract_xlsx_tables(DEEP_BANNER)
    out["deep_headers"] = [t["headers"] for t in tables_deep]

    tables_ranked = ingest.extract_xlsx_tables(RANKED)
    out["ranked_headers"] = [t["headers"] for t in tables_ranked]
    out["ranked_numeric_cols"] = [t["numeric_columns"] for t in tables_ranked]

    return out


# Captured from the UNMODIFIED extractors, immediately before the P1-14
# banner fix. Ordinary spreadsheets must hash to these forever; if either
# changes, chunk text (and therefore embeddings) changed for every future
# spreadsheet ingest, which is never allowed to happen by accident.
BASELINE_NORMAL_PROSE = "f624c69d88c0fa73c12b9748a0f428ae29e506c6177f9eaa09cde621de3aaeff"
BASELINE_NORMAL_TABLES = "4f7773bbc684a79564daeaf1722884db775ce99fb033bdbd0a672c44d906a681"


def verify() -> list[str]:
    r = report()
    fails = []

    # 1. THE REGRESSION GUARD: ordinary sheets untouched, byte for byte.
    if r["normal_prose_sha256"] != BASELINE_NORMAL_PROSE:
        fails.append(f"extract_xlsx CHANGED for a normal sheet: {r['normal_prose_sha256']}")
    if r["normal_tables_sha256"] != BASELINE_NORMAL_TABLES:
        fails.append(f"extract_xlsx_tables CHANGED for a normal sheet: {r['normal_tables_sha256']}")

    # 2. THE FIX: a banner sheet resolves to the real column names.
    bh = r["banner_headers"][0]
    if bh != ["Department", "Q1 Actual", "Q2 Actual", "Variance"]:
        fails.append(f"banner headers not resolved: {bh}")
    if any(str(h).startswith("col_") for h in bh):
        fails.append(f"banner headers still contain placeholders: {bh}")
    if r["banner_first_row"][0].get("Department") != "Engineering":
        fails.append(f"banner first data row wrong: {r['banner_first_row'][0]}")
    if r["banner_numeric_cols"][0] != ["Q1 Actual", "Q2 Actual", "Variance"]:
        fails.append(f"banner numeric_columns not meaningful: {r['banner_numeric_cols'][0]}")

    # 3. THE SAFETY BOUND: a real one-column sheet keeps its own header.
    if r["onecol_headers"][0] != ["Note"] or r["onecol_row_count"][0] != 2:
        fails.append(f"one-column sheet was damaged: {r['onecol_headers']} {r['onecol_row_count']}")

    # 4. A two-row banner (title + subtitle) before the real header resolves
    #    correctly — matches the shape of live "Executive Dashboard" sheets.
    twh = r["tworow_headers"][0]
    if twh != ["Metric", "Target", "Actual", "Variance", "Status"]:
        fails.append(f"two-row banner not resolved: {twh}")
    if r["tworow_numeric_cols"][0] != ["Target", "Actual", "Variance"]:
        fails.append(f"two-row banner numeric_columns not meaningful: {r['tworow_numeric_cols'][0]}")

    # 5. A banner deeper than _BANNER_SCAN_LIMIT is left alone, not
    #    partially/incorrectly resolved — proves the bound is real, not
    #    just untested.
    dh = r["deep_headers"][0]
    if dh[0] != "TITLE":
        fails.append(f"deep banner should have been left alone, got: {dh}")

    # 6. A numeric-looking rank/ID column right under a real header is never
    #    mistaken for evidence the row above it is data, not header — header
    #    detection must key off ROW WIDTH only, never cell content.
    rh = r["ranked_headers"][0]
    if rh != ["Rank", "Rep", "Deals Closed", "Revenue"]:
        fails.append(f"ranked-column sheet not resolved: {rh}")
    if r["ranked_numeric_cols"][0] != ["Rank", "Deals Closed", "Revenue"]:
        fails.append(f"ranked-column numeric_columns wrong: {r['ranked_numeric_cols'][0]}")

    return fails


if __name__ == "__main__":
    r = report()
    for k, v in r.items():
        print(f"{k}: {v}")
    print()
    problems = verify()
    if problems:
        for p in problems:
            print(f"FAIL: {p}")
        raise SystemExit(1)
    print("ALL CHECKS PASSED "
          "(normal sheets byte-identical, banner resolved, one-column safe)")
